"""End-to-end pipeline execution and verification test suite."""

import os
from pathlib import Path
import sqlite3
import subprocess
import sys
from typing import Generator

import openpyxl
import pypdf
import pytest
from typer.testing import CliRunner

from src.cli.main import (
    PipelineOrchestrator,
    PipelineSummary,
    app,
    run_pipeline_verification,
)

runner = CliRunner()


@pytest.fixture
def raw_data_path() -> Path:
    """Return path to project raw dataset directory."""
    return Path(__file__).resolve().parent.parent.parent / "data" / "raw"


@pytest.fixture
def verify_script_path() -> Path:
    """Return path to standalone verify_pipeline.py script."""
    return Path(__file__).resolve().parent.parent.parent / "scripts" / "verify_pipeline.py"


# --- 1. End-to-End CLI Pipeline Execution Tests ---

class TestFullPipelineCLI:
    """Validate end-to-end execution of full pipeline through CLI interface."""

    def test_full_pipeline_cli_execution_and_artifacts(
        self,
        raw_data_path: Path,
        tmp_path: Path,
    ) -> None:
        """Verify CLI run-all executes all stages and produces valid deliverables."""
        target_db = tmp_path / "e2e_cli_test.db"
        output_reports = tmp_path / "e2e_cli_reports"

        # 1. Execute run-all via Typer CliRunner
        result = runner.invoke(
            app,
            [
                "run-all",
                "--data-dir",
                str(raw_data_path),
                "--db-path",
                str(target_db),
                "--output-dir",
                str(output_reports),
            ],
        )

        assert result.exit_code == 0, f"CLI execution failed with stdout:\n{result.stdout}"
        assert "PIPELINE EXECUTION COMPLETED SUCCESSFULLY" in result.stdout
        assert "1. ETL Ingestion" in result.stdout
        assert "2. Semantic Matching" in result.stdout
        assert "3. BI Reporting" in result.stdout

        # 2. Verify Database Existence and 10 Tables
        assert target_db.exists()
        assert target_db.stat().st_size > 0

        con = sqlite3.connect(str(target_db))
        try:
            expected_tables = {
                "skills",
                "candidates",
                "candidate_skills",
                "candidate_experiences",
                "job_descriptions",
                "job_skills",
                "matching_evaluations",
                "operational_metrics",
                "ingestion_batches",
                "ingestion_errors",
            }
            rows = con.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()
            found = {r[0] for r in rows if not r[0].startswith("sqlite_")}
            assert expected_tables.issubset(found), f"Missing tables: {expected_tables - found}"

            # 3. Verify Minimum Record Counts
            cand_count = con.execute("SELECT count(*) FROM candidates").fetchone()[0]
            job_count = con.execute("SELECT count(*) FROM job_descriptions").fetchone()[0]
            metric_count = con.execute("SELECT count(*) FROM operational_metrics").fetchone()[0]
            eval_count = con.execute("SELECT count(*) FROM matching_evaluations").fetchone()[0]

            assert cand_count >= 10, f"Expected >= 10 candidates, found {cand_count}"
            assert job_count >= 4, f"Expected >= 4 jobs, found {job_count}"
            assert metric_count >= 36, f"Expected >= 36 operational metrics, found {metric_count}"
            assert eval_count >= 40, f"Expected >= 40 evaluations, found {eval_count}"

            # 4. Verify Foreign Key Referential Integrity
            fk_violations = con.execute("PRAGMA foreign_key_check").fetchall()
            assert len(fk_violations) == 0, f"Foreign key violations detected: {fk_violations}"

            # 5. Verify Score Bounds
            score_bounds = con.execute(
                """
                SELECT
                    min(overall_score), max(overall_score),
                    min(skills_score), max(skills_score),
                    min(experience_score), max(experience_score),
                    min(semantic_similarity), max(semantic_similarity)
                FROM matching_evaluations
                """
            ).fetchone()

            min_ov, max_ov = score_bounds[0], score_bounds[1]
            min_sk, max_sk = score_bounds[2], score_bounds[3]
            min_ex, max_ex = score_bounds[4], score_bounds[5]
            min_sim, max_sim = score_bounds[6], score_bounds[7]

            assert 0.0 <= min_ov <= 100.0 and 0.0 <= max_ov <= 100.0
            assert 0.0 <= min_sk <= 100.0 and 0.0 <= max_sk <= 100.0
            assert 0.0 <= min_ex <= 100.0 and 0.0 <= max_ex <= 100.0
            assert 0.0 <= min_sim <= 1.0 and 0.0 <= max_sim <= 1.0
        finally:
            con.close()

        # 6. Verify Excel Deliverable
        excel_file = output_reports / "talent_bi_report.xlsx"
        assert excel_file.exists()
        assert excel_file.stat().st_size > 5000

        wb = openpyxl.load_workbook(str(excel_file), data_only=True)
        expected_sheets = [
            "Executive Dashboard",
            "Candidates Ranking",
            "Gap Analysis",
            "Operational Metrics",
        ]
        assert wb.sheetnames == expected_sheets

        for s_name in expected_sheets:
            ws = wb[s_name]
            assert ws.max_row >= 2, f"Sheet {s_name} has insufficient rows: {ws.max_row}"
            assert ws.max_column >= 2, f"Sheet {s_name} has insufficient columns: {ws.max_column}"

        # 7. Verify PDF Deliverable
        pdf_file = output_reports / "executive_evaluation_summary.pdf"
        assert pdf_file.exists()
        assert pdf_file.stat().st_size > 5120

        with open(pdf_file, "rb") as pf:
            header = pf.read(5)
            assert header.startswith(b"%PDF-")

        pdf_reader = pypdf.PdfReader(str(pdf_file))
        assert len(pdf_reader.pages) >= 1
        total_text = "".join(p.extract_text() or "" for p in pdf_reader.pages)
        assert len(total_text) > 100
        assert "SMART TALENT BI" in total_text

        # 8. Run CLI verify Subcommand on Generated Output
        verify_res = runner.invoke(
            app,
            [
                "verify",
                "--db-path",
                str(target_db),
                "--reports-dir",
                str(output_reports),
            ],
        )
        assert verify_res.exit_code == 0
        assert "14/14 checks passed" in verify_res.stdout
        assert "Status: PASS" in verify_res.stdout


# --- 2. Programmatic End-to-End Orchestrator Tests ---

class TestFullPipelineProgrammatic:
    """Validate PipelineOrchestrator programmatic entry point."""

    def test_orchestrator_run_all_contract(
        self,
        raw_data_path: Path,
        tmp_path: Path,
    ) -> None:
        """Verify run_all programmatic method adheres to interface contract."""
        target_db = tmp_path / "orchestrator_test.db"
        output_reports = tmp_path / "orchestrator_reports"

        orchestrator = PipelineOrchestrator()
        summary: PipelineSummary = orchestrator.run_all(
            data_dir=raw_data_path,
            db_path=target_db,
            output_dir=output_reports,
        )

        assert summary.success is True
        assert summary.total_duration_seconds > 0.0

        # Ingestion Result Check
        assert summary.ingestion.success is True
        assert summary.ingestion.total_loaded >= 49
        assert summary.ingestion.total_rejected >= 0
        assert len(summary.ingestion.batch_summaries) == 3

        # Matching Result Check
        assert summary.matching.success is True
        assert summary.matching.evaluations_count == 40
        assert summary.matching.average_score > 0.0
        assert len(summary.matching.evaluations) == 40

        # Reporting Result Check
        assert summary.reporting.success is True
        assert summary.reporting.excel_path.exists()
        assert summary.reporting.pdf_path.exists()
        assert summary.reporting.kpi_summary is not None
        assert summary.reporting.kpi_summary.total_candidates == 10
        assert summary.reporting.kpi_summary.total_jobs == 4
        assert summary.reporting.kpi_summary.total_evaluations == 40

        # Run verification routine
        checks, passed = run_pipeline_verification(db_path=target_db, reports_dir=output_reports)
        assert passed is True
        assert len(checks) == 14
        assert all(c.passed for c in checks)


# --- 3. Standalone Verification Script Tests ---

class TestStandaloneVerificationScript:
    """Verify scripts/verify_pipeline.py standalone execution."""

    def test_standalone_script_success(
        self,
        raw_data_path: Path,
        verify_script_path: Path,
        tmp_path: Path,
    ) -> None:
        """Verify scripts/verify_pipeline.py exits 0 and displays structured table."""
        target_db = tmp_path / "script_test.db"
        output_reports = tmp_path / "script_reports"

        orchestrator = PipelineOrchestrator()
        orchestrator.run_all(
            data_dir=raw_data_path,
            db_path=target_db,
            output_dir=output_reports,
        )

        cmd = [
            sys.executable,
            str(verify_script_path),
            "--db-path",
            str(target_db),
            "--reports-dir",
            str(output_reports),
        ]

        proc = subprocess.run(cmd, capture_output=True, text=True)
        assert proc.returncode == 0, f"Script failed with output:\n{proc.stderr}\n{proc.stdout}"
        assert "14/14 checks passed" in proc.stdout
        assert "Verdict: PASSED" in proc.stdout

    def test_standalone_script_failure_on_missing_db(
        self,
        verify_script_path: Path,
        tmp_path: Path,
    ) -> None:
        """Verify scripts/verify_pipeline.py exits 1 on non-existent database."""
        missing_db = tmp_path / "missing_db.db"
        missing_reports = tmp_path / "missing_reports"

        cmd = [
            sys.executable,
            str(verify_script_path),
            "--db-path",
            str(missing_db),
            "--reports-dir",
            str(missing_reports),
        ]

        proc = subprocess.run(cmd, capture_output=True, text=True)
        assert proc.returncode == 1
        assert "FAIL" in proc.stdout or "FAILED" in proc.stdout


# --- 4. Cleanliness and Policy Compliance Tests ---

class TestCleanlinessAndComplianceE2E:
    """Enforce encoding cleanliness in e2e test suite."""

    def test_zero_emojis_in_e2e_codebase(self) -> None:
        """Verify e2e test files contain zero emojis."""
        e2e_dir = Path(__file__).resolve().parent
        for py_file in e2e_dir.glob("*.py"):
            text = py_file.read_text(encoding="utf-8")
            for char in text:
                cp = ord(char)
                assert not (
                    0x1F600 <= cp <= 0x1F64F
                    or 0x1F300 <= cp <= 0x1F5FF
                    or 0x1F900 <= cp <= 0x1F9FF
                    or 0x2600 <= cp <= 0x27BF
                    or 0x1F000 <= cp <= 0x1FFFF
                ), f"Prohibited emoji '{char}' detected in {py_file}"

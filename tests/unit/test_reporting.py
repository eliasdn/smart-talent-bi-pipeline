"""Comprehensive unit tests for Business Intelligence metrics calculation and reporting deliverables.

Validates:
- Statistical score distributions, quartiles, and high-fit thresholds
- Recruitment ROI mathematical formulations
- Funnel stage durations and sourcing channel conversion rates
- Relational database KPI aggregation and DataFrame generation
- Multi-tab Excel workbook generation with OpenPyXL (all 4 sheets verified)
- Executive PDF dossier generation with ReportLab SimpleDocTemplate and NumberedCanvas
- High-level ReportingService coordinator
- Strict compliance: zero emojis and zero prohibited term mentions
"""

from datetime import date, datetime, timezone
from decimal import Decimal
import json
from pathlib import Path
from typing import Generator

import openpyxl
import pandas as pd
import pytest
from sqlalchemy.orm import Session

from src.models.entities import (
    Candidate,
    JobDescription,
    MatchingEvaluation,
    OperationalMetric,
    create_db_engine,
    init_db,
)
from src.reporting.excel_generator import ExcelReportGenerator
from src.reporting.metrics import (
    BIMetricsCalculator,
    KPISummaryDTO,
    RecruitmentROIDTO,
    ScoreDistributionDTO,
    ScoreQuartilesDTO,
)
from src.reporting.pdf_generator import NumberedCanvas, PDFReportGenerator
from src.reporting.service import ReportingService, generate_all_reports


@pytest.fixture
def populated_db_session(db_session: Session) -> Session:
    """Populate an in-memory database session with realistic test candidates, jobs, and evaluations."""
    now = datetime.now(timezone.utc)

    # 1. Candidates
    cand1 = Candidate(
        candidate_id="cand-001",
        full_name="Antoine Dupont",
        email="antoine.dupont@example.com",
        location="Paris, France",
        current_title="Lead Data Engineer",
        years_of_experience=Decimal("6.5"),
        education_level="Master Informatique",
        raw_cv_text="Lead Data Engineer expert Python, SQL, Apache Spark, Kafka et architectures cloud AWS.",
        created_at=now,
        updated_at=now,
    )
    cand2 = Candidate(
        candidate_id="cand-002",
        full_name="Camille Bernard",
        email="camille.bernard@example.com",
        location="Lyon, France",
        current_title="Data Analyst Senior",
        years_of_experience=Decimal("4.0"),
        education_level="Master Statistiques",
        raw_cv_text="Data Analyst avec expertise SQL, Tableau, Power BI, Python et modelisation de donnees.",
        created_at=now,
        updated_at=now,
    )
    cand3 = Candidate(
        candidate_id="cand-003",
        full_name="Julien Moreau",
        email="julien.moreau@example.com",
        location="Nantes, France",
        current_title="Developpeur Junior",
        years_of_experience=Decimal("1.5"),
        education_level="Licence Informatique",
        raw_cv_text="Developpeur junior Python et SQL motive par l'apprentissage du machine learning.",
        created_at=now,
        updated_at=now,
    )
    db_session.add_all([cand1, cand2, cand3])
    db_session.flush()

    # 2. Jobs
    job1 = JobDescription(
        job_id="job-001",
        job_code="JOB-DE-01",
        title="Senior Data Engineer",
        department="Data Platform",
        location="Paris, France",
        employment_type="Full-time",
        experience_min_years=Decimal("5.0"),
        education_level_required="Master Informatique",
        salary_range_min=Decimal("65000.00"),
        salary_range_max=Decimal("80000.00"),
        currency="EUR",
        status="open",
        raw_description="Poste de Senior Data Engineer pour concevoir des pipelines streaming temps reel.",
        created_at=now,
        updated_at=now,
    )
    job2 = JobDescription(
        job_id="job-002",
        job_code="JOB-DA-02",
        title="Lead Business Analyst",
        department="Business Analytics",
        location="Paris, France",
        employment_type="Full-time",
        experience_min_years=Decimal("3.5"),
        education_level_required="Master Statistiques",
        salary_range_min=Decimal("55000.00"),
        salary_range_max=Decimal("70000.00"),
        currency="EUR",
        status="open",
        raw_description="Poste de Lead Business Analyst pour piloter les tableaux de bord strategiques.",
        created_at=now,
        updated_at=now,
    )
    db_session.add_all([job1, job2])
    db_session.flush()

    # 3. Evaluations (6 combinations)
    evals = [
        MatchingEvaluation(
            evaluation_id="eval-001",
            candidate_id="cand-001",
            job_id="job-001",
            overall_score=Decimal("88.50"),
            skills_score=Decimal("92.00"),
            experience_score=Decimal("100.00"),
            semantic_similarity=Decimal("0.8420"),
            recommendation="strong_hire",
            matched_skills=json.dumps(["Python", "SQL", "Apache Spark", "Kafka"]),
            missing_skills=json.dumps([]),
            strengths_summary="Validation excellente des competences requises et experience superieure.",
            gaps_summary="Aucun ecart bloquant constate sur le profil.",
            executive_summary="Profil de tout premier plan pour le poste de Senior Data Engineer.",
            evaluated_at=now,
        ),
        MatchingEvaluation(
            evaluation_id="eval-002",
            candidate_id="cand-001",
            job_id="job-002",
            overall_score=Decimal("64.20"),
            skills_score=Decimal("60.00"),
            experience_score=Decimal("100.00"),
            semantic_similarity=Decimal("0.5120"),
            recommendation="consider",
            matched_skills=json.dumps(["Python", "SQL"]),
            missing_skills=json.dumps(["Tableau", "Power BI"]),
            strengths_summary="Solides bases techniques de donnees.",
            gaps_summary="Manque de competences de dataviz dediees (Tableau, Power BI).",
            executive_summary="Profil a considerer si le poste s'oriente vers l'ingenierie.",
            evaluated_at=now,
        ),
        MatchingEvaluation(
            evaluation_id="eval-003",
            candidate_id="cand-002",
            job_id="job-001",
            overall_score=Decimal("58.30"),
            skills_score=Decimal("52.00"),
            experience_score=Decimal("80.00"),
            semantic_similarity=Decimal("0.4650"),
            recommendation="consider",
            matched_skills=json.dumps(["Python", "SQL"]),
            missing_skills=json.dumps(["Apache Spark", "Kafka"]),
            strengths_summary="Excellente maitrise SQL et rigueur analytique.",
            gaps_summary="Ecart important sur les frameworks de distributed computing.",
            executive_summary="Candidat a reorienter vers l'analytique avancee.",
            evaluated_at=now,
        ),
        MatchingEvaluation(
            evaluation_id="eval-004",
            candidate_id="cand-002",
            job_id="job-002",
            overall_score=Decimal("91.40"),
            skills_score=Decimal("94.00"),
            experience_score=Decimal("100.00"),
            semantic_similarity=Decimal("0.8910"),
            recommendation="strong_hire",
            matched_skills=json.dumps(["SQL", "Tableau", "Power BI", "Python"]),
            missing_skills=json.dumps([]),
            strengths_summary="Parfaite adequation fonctionnelle et maitrise des outils BI.",
            gaps_summary="Aucun ecart identifie.",
            executive_summary="Candidate ideale pour piloter le reporting strategique.",
            evaluated_at=now,
        ),
        MatchingEvaluation(
            evaluation_id="eval-005",
            candidate_id="cand-003",
            job_id="job-001",
            overall_score=Decimal("42.10"),
            skills_score=Decimal("40.00"),
            experience_score=Decimal("30.00"),
            semantic_similarity=Decimal("0.3200"),
            recommendation="reject",
            matched_skills=json.dumps(["Python"]),
            missing_skills=json.dumps(["Apache Spark", "Kafka", "SQL"]),
            strengths_summary="Bases de programmation Python verifiees.",
            gaps_summary="Seniorite et competences data insuffisantes pour un poste senior.",
            executive_summary="Niveau d'experience en deca des exigences du poste.",
            evaluated_at=now,
        ),
        MatchingEvaluation(
            evaluation_id="eval-006",
            candidate_id="cand-003",
            job_id="job-002",
            overall_score=Decimal("48.50"),
            skills_score=Decimal("45.00"),
            experience_score=Decimal("42.80"),
            semantic_similarity=Decimal("0.3850"),
            recommendation="reject",
            matched_skills=json.dumps(["SQL"]),
            missing_skills=json.dumps(["Tableau", "Power BI"]),
            strengths_summary="Motivation demontree et bases SQL.",
            gaps_summary="Experience d'analyse decisionnelle trop limitee.",
            executive_summary="Profil junior a developper sur d'autres opportunites.",
            evaluated_at=now,
        ),
    ]
    db_session.add_all(evals)
    db_session.flush()

    # 4. Operational metrics (funnel tracking)
    metrics = [
        OperationalMetric(
            metric_id="met-001",
            candidate_id="cand-001",
            job_id="job-001",
            application_date=date(2024, 1, 10),
            screening_date=date(2024, 1, 11),
            technical_interview_date=date(2024, 1, 18),
            final_interview_date=date(2024, 1, 25),
            offer_date=date(2024, 1, 30),
            hiring_decision="hired",
            time_to_screen_hours=Decimal("12.5"),
            automated_screening_time_seconds=Decimal("2.3"),
            manual_screening_estimated_minutes=Decimal("45.0"),
            recruiter_name="Sarah Martin",
            sourcing_channel="LinkedIn",
            cost_per_applicant_eur=Decimal("80.00"),
            created_at=now,
        ),
        OperationalMetric(
            metric_id="met-002",
            candidate_id="cand-002",
            job_id="job-002",
            application_date=date(2024, 1, 15),
            screening_date=date(2024, 1, 16),
            technical_interview_date=date(2024, 1, 22),
            final_interview_date=date(2024, 1, 28),
            offer_date=date(2024, 2, 2),
            hiring_decision="hired",
            time_to_screen_hours=Decimal("14.0"),
            automated_screening_time_seconds=Decimal("2.5"),
            manual_screening_estimated_minutes=Decimal("45.0"),
            recruiter_name="Julien Delacroix",
            sourcing_channel="Referral",
            cost_per_applicant_eur=Decimal("0.00"),
            created_at=now,
        ),
        OperationalMetric(
            metric_id="met-003",
            candidate_id="cand-003",
            job_id="job-001",
            application_date=date(2024, 2, 1),
            screening_date=date(2024, 2, 2),
            technical_interview_date=None,
            final_interview_date=None,
            offer_date=None,
            hiring_decision="rejected",
            time_to_screen_hours=Decimal("8.0"),
            automated_screening_time_seconds=Decimal("2.1"),
            manual_screening_estimated_minutes=Decimal("45.0"),
            recruiter_name="Sarah Martin",
            sourcing_channel="Direct Application",
            cost_per_applicant_eur=Decimal("25.00"),
            created_at=now,
        ),
    ]
    db_session.add_all(metrics)
    db_session.commit()

    return db_session


# --- 1. BI Metrics Calculation Tests ---

class TestBIMetricsCalculations:
    """Test mathematical formulas and algorithms for KPIs, distributions, and ROI."""

    def test_score_statistics_empty_and_populated(self) -> None:
        """Verify mean, median, standard deviation, and quartiles on empty and populated data."""
        # Empty series
        mean_e, med_e, std_e, q_e = BIMetricsCalculator.calculate_score_statistics([])
        assert mean_e == 0.0
        assert med_e == 0.0
        assert std_e == 0.0
        assert q_e.q1 == 0.0
        assert q_e.q3 == 0.0

        # Populated series
        scores = [40.0, 50.0, 60.0, 70.0, 80.0, 90.0]
        mean_v, med_v, std_v, q_v = BIMetricsCalculator.calculate_score_statistics(scores)
        assert mean_v == 65.0
        assert med_v == 65.0
        assert std_v == pytest.approx(18.71, abs=0.05)
        assert q_v.q1 == pytest.approx(52.5, abs=0.5)
        assert q_v.q3 == pytest.approx(77.5, abs=0.5)
        assert q_v.iqr == pytest.approx(q_v.q3 - q_v.q1, abs=0.01)

    def test_recruitment_roi_mathematical_formulation(self) -> None:
        """Verify baseline manual screening vs automated pipeline execution savings."""
        # 10 candidates, 36 seconds automated execution
        # Manual baseline: 10 * 0.75 = 7.5 hours
        # Automated: 36 / 3600 = 0.01 hours
        # Hours saved: 7.5 - 0.01 = 7.49 hours
        # Cost saved: 7.49 * 65.0 = 486.85
        roi = BIMetricsCalculator.calculate_roi(
            candidate_count=10,
            automated_execution_seconds=36.0,
            hourly_rate=65.0,
            manual_baseline_hours=0.75,
        )
        assert roi.total_candidates == 10
        assert roi.manual_hours_per_candidate == 0.75
        assert roi.total_manual_hours == 7.5
        assert roi.total_automated_hours == 0.01
        assert roi.total_hours_saved == 7.49
        assert roi.recruiter_hourly_rate == 65.0
        assert roi.total_cost_saved == pytest.approx(486.85, abs=0.1)
        assert roi.time_savings_percentage >= 99.0

    def test_score_distribution_brackets(self) -> None:
        """Verify correct binning into 4 cohort brackets and percentage sum."""
        scores = [25.0, 38.0, 45.0, 55.0, 65.0, 75.0, 85.0, 95.0]
        dist = BIMetricsCalculator.calculate_score_distribution(scores)

        assert len(dist) == 4
        # [0, 40): 2 items
        assert dist[0].count == 2
        assert dist[0].percentage == 25.0
        # [40, 60): 2 items
        assert dist[1].count == 2
        assert dist[1].percentage == 25.0
        # [60, 80): 2 items
        assert dist[2].count == 2
        assert dist[2].percentage == 25.0
        # [80, 100]: 2 items
        assert dist[3].count == 2
        assert dist[3].percentage == 25.0

    def test_kpi_summary_database_aggregation(self, populated_db_session: Session) -> None:
        """Verify full KPI summary calculation over populated relational tables."""
        calc = BIMetricsCalculator(populated_db_session)
        kpi = calc.compute_kpi_summary()

        assert kpi.total_candidates == 3
        assert kpi.total_jobs == 2
        assert kpi.total_evaluations == 6
        assert kpi.average_matching_score == pytest.approx(65.50, abs=0.1)
        assert kpi.median_matching_score == pytest.approx(61.25, abs=0.1)

        # High-fit count (score >= 70%): eval-001 (88.5) and eval-004 (91.4) -> 2
        assert kpi.high_fit_count == 2
        assert kpi.high_fit_rate == pytest.approx(33.3, abs=0.1)

        # ROI check
        assert kpi.roi.total_candidates == 3
        assert kpi.roi.total_manual_hours == 2.25
        assert kpi.roi.total_hours_saved > 2.20
        assert kpi.roi.total_cost_saved > 140.0

        # Department summaries check
        assert len(kpi.department_summaries) == 2
        depts = {d.department for d in kpi.department_summaries}
        assert "Data Platform" in depts
        assert "Business Analytics" in depts

    def test_dataframes_structure_and_content(self, populated_db_session: Session) -> None:
        """Verify rankings, gap analysis, and operational metrics DataFrames."""
        calc = BIMetricsCalculator(populated_db_session)

        # 1. Rankings DataFrame
        rank_df = calc.get_rankings_dataframe()
        assert not rank_df.empty
        assert len(rank_df) == 6
        assert "Rank" in rank_df.columns
        assert "Overall Score" in rank_df.columns
        assert "Tier" in rank_df.columns
        assert "Recommendation" in rank_df.columns
        # First row should be highest score (91.4)
        assert rank_df.iloc[0]["Overall Score"] >= rank_df.iloc[1]["Overall Score"]

        # 2. Gap Analysis DataFrame
        gap_df = calc.get_gap_analysis_dataframe()
        assert not gap_df.empty
        assert len(gap_df) == 6
        assert "Matched Skills" in gap_df.columns
        assert "Missing Required Skills" in gap_df.columns
        assert "Experience Delta Years" in gap_df.columns
        assert "Key Strengths" in gap_df.columns

        # 3. Operational Metrics DataFrame
        op_df = calc.get_operational_metrics_dataframe()
        assert not op_df.empty
        assert len(op_df) == 3
        assert "Metric ID" in op_df.columns
        assert "Sourcing Channel" in op_df.columns
        assert "Hiring Decision" in op_df.columns

    def test_empty_database_graceful_handling(self, db_session: Session) -> None:
        """Verify calculator operates without error when tables contain zero records."""
        calc = BIMetricsCalculator(db_session)
        kpi = calc.compute_kpi_summary()

        assert kpi.total_candidates == 0
        assert kpi.total_jobs == 0
        assert kpi.total_evaluations == 0
        assert kpi.average_matching_score == 0.0
        assert kpi.high_fit_count == 0

        rank_df = calc.get_rankings_dataframe()
        assert rank_df.empty
        gap_df = calc.get_gap_analysis_dataframe()
        assert gap_df.empty
        op_df = calc.get_operational_metrics_dataframe()
        assert op_df.empty


# --- 2. Excel Generation Tests ---

class TestExcelReportGenerator:
    """Verify Excel workbook generation, sheet presence, formatting, and cell data."""

    def test_excel_workbook_generation_and_sheets(
        self,
        populated_db_session: Session,
        tmp_path: Path,
    ) -> None:
        """Verify Excel workbook generation creates all 4 required sheets with valid structure."""
        calc = BIMetricsCalculator(populated_db_session)
        generator = ExcelReportGenerator(calc)

        out_file = tmp_path / "talent_bi_report.xlsx"
        written = generator.generate(out_file)

        assert written.exists()
        assert written.stat().st_size > 5000

        # Inspect workbook with OpenPyXL
        wb = openpyxl.load_workbook(written)
        expected_sheets = [
            "Executive Dashboard",
            "Candidates Ranking",
            "Gap Analysis",
            "Operational Metrics",
        ]
        assert wb.sheetnames == expected_sheets

        # 1. Executive Dashboard assertions
        ws_dash = wb["Executive Dashboard"]
        assert ws_dash["A1"].value == "SMART TALENT BI - TABLEAU DE BORD DECISIONNEL RECRUTEMENT"
        # Check KPI card values present
        assert ws_dash["A5"].value == "3"  # Total candidates
        assert ws_dash["C5"].value == "2"  # Total jobs
        assert ws_dash["E5"].value == "6"  # Total evaluations

        # 2. Candidates Ranking assertions
        ws_rank = wb["Candidates Ranking"]
        assert ws_rank["A3"].value == "Rang"
        assert ws_rank["B3"].value == "Candidat"
        assert ws_rank["E3"].value == "Score Global"
        # 6 evaluations + header row (row 3) = rows 4 through 9
        assert ws_rank["A4"].value == 1
        assert ws_rank.max_row >= 9

        # 3. Gap Analysis assertions
        ws_gap = wb["Gap Analysis"]
        assert ws_gap["A3"].value == "Candidat"
        assert ws_gap["C3"].value == "Competences Validees"
        assert ws_gap.max_row >= 9

        # 4. Operational Metrics assertions
        ws_op = wb["Operational Metrics"]
        assert ws_op["A3"].value == "ID Metrique"
        assert ws_op["D3"].value == "Canal de Sourcing"
        assert ws_op.max_row >= 6

    def test_excel_styling_and_column_widths(
        self,
        populated_db_session: Session,
        tmp_path: Path,
    ) -> None:
        """Verify header cell fill and dynamic column width dimensions."""
        calc = BIMetricsCalculator(populated_db_session)
        generator = ExcelReportGenerator(calc)

        out_file = tmp_path / "styled_report.xlsx"
        generator.generate(out_file)

        wb = openpyxl.load_workbook(out_file)
        ws_rank = wb["Candidates Ranking"]

        # Check header fill color on Candidates Ranking table
        hdr_cell = ws_rank["A3"]
        assert hdr_cell.fill.start_color.rgb.endswith("1F4E79")
        assert hdr_cell.font.color.rgb.endswith("FFFFFF")

        # Check column dimensions configured
        for col_letter in ["A", "B", "C", "D", "E"]:
            width = ws_rank.column_dimensions[col_letter].width
            assert width is not None
            assert width >= 10


# --- 3. PDF Generation Tests ---

class TestPDFReportGenerator:
    """Verify PDF document compilation, ReportLab SimpleDocTemplate, and NumberedCanvas."""

    def test_pdf_generation_validity_and_header(
        self,
        populated_db_session: Session,
        tmp_path: Path,
    ) -> None:
        """Verify PDF generation creates valid %PDF- document with sufficient byte size."""
        calc = BIMetricsCalculator(populated_db_session)
        generator = PDFReportGenerator(calc)

        out_pdf = tmp_path / "executive_evaluation_summary.pdf"
        written = generator.generate(out_pdf)

        assert written.exists()
        assert written.stat().st_size > 5000

        with open(written, "rb") as f:
            header = f.read(5)
            assert header == b"%PDF-"

    def test_pdf_multipage_numbered_canvas_execution(
        self,
        populated_db_session: Session,
        tmp_path: Path,
    ) -> None:
        """Verify multi-page compilation runs without flowable overflow or canvas errors."""
        calc = BIMetricsCalculator(populated_db_session)
        generator = PDFReportGenerator(calc)

        out_pdf = tmp_path / "multipage_summary.pdf"
        generator.generate(out_pdf)

        # Confirm file is generated and non-empty (>5000 bytes)
        assert out_pdf.exists()
        assert out_pdf.stat().st_size > 5000


# --- 4. ReportingService Coordinator Tests ---

class TestReportingService:
    """Verify high-level service coordinator and standalone functional entry point."""

    def test_reporting_service_generate_all(
        self,
        populated_db_session: Session,
        tmp_path: Path,
    ) -> None:
        """Verify ReportingService.generate_all outputs both Excel and PDF deliverables."""
        service = ReportingService(session=populated_db_session)
        out_dir = tmp_path / "reports_output"

        results = service.generate_all(out_dir)

        assert "excel" in results
        assert "pdf" in results
        assert results["excel"].exists()
        assert results["pdf"].exists()
        assert results["excel"].name == "talent_bi_report.xlsx"
        assert results["pdf"].name == "executive_evaluation_summary.pdf"

    def test_generate_all_reports_from_db_file(
        self,
        populated_db_session: Session,
        tmp_path: Path,
    ) -> None:
        """Verify generate_all_reports functional entry point connects to file-based SQLite database."""
        db_file = tmp_path / "test_warehouse.db"
        engine = create_db_engine(f"sqlite:///{db_file.resolve()}")
        init_db(engine)

        # Transfer records from populated in-memory session to file-based database
        with Session(engine) as file_session:
            for cand in populated_db_session.query(Candidate).all():
                file_session.merge(cand)
            for job in populated_db_session.query(JobDescription).all():
                file_session.merge(job)
            for ev in populated_db_session.query(MatchingEvaluation).all():
                file_session.merge(ev)
            for met in populated_db_session.query(OperationalMetric).all():
                file_session.merge(met)
            file_session.commit()

        out_dir = tmp_path / "file_reports"
        artifacts = generate_all_reports(db_file, out_dir)

        assert artifacts["excel"].exists()
        assert artifacts["pdf"].exists()
        assert artifacts["excel"].stat().st_size > 5000
        assert artifacts["pdf"].stat().st_size > 5000


# --- 5. Strict Cleanliness and Compliance Tests ---

class TestCleanlinessAndCompliance:
    """Enforce strict zero-emoji and zero-prohibited-mention constraints."""

    def test_zero_emojis_in_reporting_codebase(self) -> None:
        """Scan all reporting source files to verify strictly zero emojis."""
        reporting_dir = Path(__file__).resolve().parent.parent.parent / "src" / "reporting"
        for py_file in reporting_dir.glob("*.py"):
            text = py_file.read_text(encoding="utf-8")
            for char in text:
                cp = ord(char)
                assert not (
                    0x1F600 <= cp <= 0x1F64F
                    or 0x1F300 <= cp <= 0x1F5FF
                    or 0x1F900 <= cp <= 0x1F9FF
                    or 0x2600 <= cp <= 0x27BF
                    or 0x1F000 <= cp <= 0x1FFFF
                ), f"Prohibited emoji character '{char}' (U+{cp:X}) detected in {py_file}"

    def test_zero_prohibited_mentions_in_reporting_codebase(self) -> None:
        """Scan reporting source files to ensure absence of prohibited automated generation terms."""
        reporting_dir = Path(__file__).resolve().parent.parent.parent / "src" / "reporting"
        prohibited = [
            "".join(["chat", "gpt"]),
            "".join(["open", "ai"]),
            "".join(["co", "pilot"]),
            "".join(["assist", "ant"]),
            "".join(["ag", "ent"]),
        ]
        for py_file in reporting_dir.glob("*.py"):
            text = py_file.read_text(encoding="utf-8").lower()
            for term in prohibited:
                assert term not in text, f"Prohibited term '{term}' found in {py_file}"

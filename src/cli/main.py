"""Talent analytics and business intelligence command-line interface."""

import os
import sys
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
import time
from typing import Any, Dict, List, Optional, Tuple, Union

# Ensure project root is available in module path for standalone script execution
_project_root = Path(__file__).resolve().parent.parent.parent
if str(_project_root) not in sys.path:
    sys.path.insert(0, str(_project_root))

import openpyxl
import pypdf
import sqlite3
import typer
from rich import box
from rich.console import Console
from rich.panel import Panel
from rich.table import Table
from sqlalchemy import select

from src.config import get_settings
from src.engine.service import MatchingEngine
from src.etl.pipeline import ETLPipeline
from src.models.entities import (
    Candidate,
    JobDescription,
    MatchingEvaluation,
    OperationalMetric,
    create_db_engine,
    get_session_factory,
    init_db,
)
from src.models.schemas import IngestionBatchSummaryDTO, MatchingEvaluationDTO
from src.reporting.metrics import BIMetricsCalculator, KPISummaryDTO
from src.reporting.service import ReportingService
from src.utils.logger import setup_logger

logger = setup_logger("smart_talent_bi.cli")
console = Console(highlight=False)

app = typer.Typer(
    name="talent-bi",
    help="Enterprise Talent Analytics & Business Intelligence CLI Pipeline.",
    no_args_is_help=True,
    add_completion=False,
)


@dataclass
class IngestionResult:
    """Outcome metrics and batch details for raw data ingestion."""

    batch_summaries: Dict[str, IngestionBatchSummaryDTO]
    total_extracted: int
    total_validated: int
    total_loaded: int
    total_rejected: int
    duration_seconds: float
    success: bool


@dataclass
class MatchingResult:
    """Outcome metrics and evaluation records from candidate-job matching."""

    evaluations_count: int
    average_score: float
    high_fit_count: int
    duration_seconds: float
    success: bool
    evaluations: List[MatchingEvaluationDTO] = field(default_factory=list)


@dataclass
class ReportingResult:
    """Outcome artifacts and key performance indicators from report generation."""

    excel_path: Path
    pdf_path: Path
    kpi_summary: Optional[KPISummaryDTO]
    duration_seconds: float
    success: bool


@dataclass
class PipelineSummary:
    """End-to-end execution summary across all pipeline stages."""

    ingestion: IngestionResult
    matching: MatchingResult
    reporting: ReportingResult
    total_duration_seconds: float
    success: bool


@dataclass
class VerificationCheckResult:
    """Status and measurement record for a single pipeline verification check."""

    check_id: str
    category: str
    description: str
    expected: str
    observed: str
    passed: bool


class PipelineOrchestrator:
    """Coordinates execution across ETL ingestion, semantic matching, and BI reporting."""

    def __init__(self, custom_logger: Optional[Any] = None) -> None:
        """Initialize orchestrator instance.

        Args:
            custom_logger: Optional logger instance.
        """
        self.logger = custom_logger or logger

    def run_ingestion(
        self,
        data_dir: Union[str, Path] = Path("data/raw"),
        db_path: Union[str, Path] = Path("data/output/talent_bi.db"),
    ) -> IngestionResult:
        """Execute raw dataset extraction, validation, and relational persistence.

        Args:
            data_dir: Directory containing raw JSON and CSV source files.
            db_path: Destination SQLite database file path.

        Returns:
            IngestionResult containing execution metrics and batch statuses.
        """
        start_time = time.perf_counter()
        raw_path = Path(data_dir).resolve()
        target_db = Path(db_path).resolve()
        target_db.parent.mkdir(parents=True, exist_ok=True)

        if not raw_path.exists():
            self.logger.error("Raw data directory does not exist: %s", raw_path)
            return IngestionResult(
                batch_summaries={},
                total_extracted=0,
                total_validated=0,
                total_loaded=0,
                total_rejected=0,
                duration_seconds=time.perf_counter() - start_time,
                success=False,
            )

        db_url = f"sqlite:///{target_db.as_posix()}"
        engine = create_db_engine(db_url)
        init_db(engine)

        etl_pipeline = ETLPipeline(engine=engine)
        summaries = etl_pipeline.run_all(raw_data_dir=raw_path)

        total_extracted = sum(b.records_extracted for b in summaries.values())
        total_validated = sum(b.records_validated for b in summaries.values())
        total_loaded = sum(b.records_loaded for b in summaries.values())
        total_rejected = sum(b.records_rejected for b in summaries.values())

        success = (
            len(summaries) > 0
            and all(b.status in ("completed", "partial") for b in summaries.values())
            and total_loaded > 0
        )
        duration = time.perf_counter() - start_time

        return IngestionResult(
            batch_summaries=summaries,
            total_extracted=total_extracted,
            total_validated=total_validated,
            total_loaded=total_loaded,
            total_rejected=total_rejected,
            duration_seconds=duration,
            success=success,
        )

    def run_matching(
        self,
        db_path: Union[str, Path] = Path("data/output/talent_bi.db"),
        persist: bool = True,
    ) -> MatchingResult:
        """Execute semantic matching between all candidates and job vacancies.

        Args:
            db_path: Path to populated SQLite database.
            persist: Whether to commit evaluation outcomes to the database.

        Returns:
            MatchingResult containing evaluation counts, score stats, and records.
        """
        start_time = time.perf_counter()
        target_db = Path(db_path).resolve()

        if not target_db.exists():
            self.logger.error("Database file does not exist: %s", target_db)
            return MatchingResult(
                evaluations_count=0,
                average_score=0.0,
                high_fit_count=0,
                duration_seconds=time.perf_counter() - start_time,
                success=False,
                evaluations=[],
            )

        if target_db.stat().st_size == 0:
            self.logger.error("Database file is empty (0 bytes): %s", target_db)
            return MatchingResult(
                evaluations_count=0,
                average_score=0.0,
                high_fit_count=0,
                duration_seconds=time.perf_counter() - start_time,
                success=False,
                evaluations=[],
            )

        try:
            db_url = f"sqlite:///{target_db.as_posix()}"
            engine = create_db_engine(db_url)
            session_factory = get_session_factory(engine)

            engine_instance = MatchingEngine()
            with session_factory() as session:
                evaluations = engine_instance.evaluate_all(session, persist=persist)

            eval_count = len(evaluations)
            avg_score = (
                float(sum(float(e.overall_score) for e in evaluations) / eval_count)
                if eval_count > 0
                else 0.0
            )
            high_fit_count = sum(1 for e in evaluations if float(e.overall_score) >= 70.0)
            duration = time.perf_counter() - start_time

            return MatchingResult(
                evaluations_count=eval_count,
                average_score=avg_score,
                high_fit_count=high_fit_count,
                duration_seconds=duration,
                success=eval_count > 0,
                evaluations=evaluations,
            )
        except Exception as exc:
            self.logger.error("Matching engine execution failed: %s", exc)
            return MatchingResult(
                evaluations_count=0,
                average_score=0.0,
                high_fit_count=0,
                duration_seconds=time.perf_counter() - start_time,
                success=False,
                evaluations=[],
            )

    def run_reporting(
        self,
        db_path: Union[str, Path] = Path("data/output/talent_bi.db"),
        output_dir: Union[str, Path] = Path("data/output/reports"),
    ) -> ReportingResult:
        """Generate executive multi-tab Excel report and PDF evaluation summary.

        Args:
            db_path: Path to populated SQLite database.
            output_dir: Target directory for document deliverables.

        Returns:
            ReportingResult containing file paths and executive KPI summary.
        """
        start_time = time.perf_counter()
        target_db = Path(db_path).resolve()
        target_out = Path(output_dir).resolve()
        target_out.mkdir(parents=True, exist_ok=True)

        if not target_db.exists():
            self.logger.error("Database file does not exist: %s", target_db)
            return ReportingResult(
                excel_path=Path(""),
                pdf_path=Path(""),
                kpi_summary=None,
                duration_seconds=time.perf_counter() - start_time,
                success=False,
            )

        if target_db.stat().st_size == 0:
            self.logger.error("Database file is empty (0 bytes): %s", target_db)
            return ReportingResult(
                excel_path=Path(""),
                pdf_path=Path(""),
                kpi_summary=None,
                duration_seconds=time.perf_counter() - start_time,
                success=False,
            )

        try:
            service = ReportingService.from_db_path(target_db)
            generated = service.generate_all(output_dir=target_out)
            kpi = service.calculator.compute_kpi_summary()
            duration = time.perf_counter() - start_time

            excel_ok = generated["excel"].exists() and generated["excel"].stat().st_size > 0
            pdf_ok = generated["pdf"].exists() and generated["pdf"].stat().st_size > 0

            return ReportingResult(
                excel_path=generated["excel"],
                pdf_path=generated["pdf"],
                kpi_summary=kpi,
                duration_seconds=duration,
                success=excel_ok and pdf_ok,
            )
        except Exception as exc:
            self.logger.error("Reporting generation failed: %s", exc)
            return ReportingResult(
                excel_path=Path(""),
                pdf_path=Path(""),
                kpi_summary=None,
                duration_seconds=time.perf_counter() - start_time,
                success=False,
            )

    def run_all(
        self,
        data_dir: Union[str, Path] = Path("data/raw"),
        db_path: Union[str, Path] = Path("data/output/talent_bi.db"),
        output_dir: Union[str, Path] = Path("data/output/reports"),
    ) -> PipelineSummary:
        """Run the complete end-to-end talent analytics pipeline.

        Args:
            data_dir: Directory containing raw input data files.
            db_path: Destination SQLite database file path.
            output_dir: Target directory for generated reports.

        Returns:
            PipelineSummary aggregating all stage results and execution times.
        """
        start_time = time.perf_counter()

        ingest_res = self.run_ingestion(data_dir=data_dir, db_path=db_path)
        if not ingest_res.success:
            total_dur = time.perf_counter() - start_time
            return PipelineSummary(
                ingestion=ingest_res,
                matching=MatchingResult(0, 0.0, 0, 0.0, False),
                reporting=ReportingResult(Path(""), Path(""), None, 0.0, False),
                total_duration_seconds=total_dur,
                success=False,
            )

        match_res = self.run_matching(db_path=db_path, persist=True)
        if not match_res.success:
            total_dur = time.perf_counter() - start_time
            return PipelineSummary(
                ingestion=ingest_res,
                matching=match_res,
                reporting=ReportingResult(Path(""), Path(""), None, 0.0, False),
                total_duration_seconds=total_dur,
                success=False,
            )

        report_res = self.run_reporting(db_path=db_path, output_dir=output_dir)
        total_dur = time.perf_counter() - start_time

        return PipelineSummary(
            ingestion=ingest_res,
            matching=match_res,
            reporting=report_res,
            total_duration_seconds=total_dur,
            success=report_res.success,
        )


def run_pipeline_verification(
    db_path: Union[str, Path],
    reports_dir: Union[str, Path],
) -> Tuple[List[VerificationCheckResult], bool]:
    """Execute comprehensive data integrity and deliverable artifact verification.

    Args:
        db_path: Path to SQLite database file.
        reports_dir: Directory containing generated reports.

    Returns:
        Tuple of (check results list, boolean indicating whether all checks passed).
    """
    target_db = Path(db_path).resolve()
    target_reports = Path(reports_dir).resolve()
    checks: List[VerificationCheckResult] = []

    # 1. Database File Existence
    if target_db.exists() and target_db.stat().st_size > 0:
        checks.append(
            VerificationCheckResult(
                check_id="DB-01",
                category="Database",
                description="Database file presence and non-zero size",
                expected="File exists and size > 0 bytes",
                observed=f"Exists, size: {target_db.stat().st_size:,} bytes",
                passed=True,
            )
        )
    else:
        checks.append(
            VerificationCheckResult(
                check_id="DB-01",
                category="Database",
                description="Database file presence and non-zero size",
                expected="File exists and size > 0 bytes",
                observed="File missing or 0 bytes",
                passed=False,
            )
        )
        return checks, False

    # Connect to SQLite for schema and record verification
    try:
        con = sqlite3.connect(str(target_db))
    except (sqlite3.DatabaseError, sqlite3.Error) as exc:
        checks.append(
            VerificationCheckResult(
                check_id="DB-02",
                category="Database",
                description="Relational 3NF schema tables presence (10 tables)",
                expected="Valid SQLite database connection",
                observed=f"Cannot connect to database: {exc}",
                passed=False,
            )
        )
        return checks, False

    try:
        # 2. Schema Tables (all 10 expected tables)
        expected_tables = {
            "skills",
            "candidates",
            "candidate_skills",
            "candidate_experiences",
            "job_descriptions",
            "job_skills",
            "matching_evaluations",
            "operational_metrics",
            "ingestion_batches",
            "ingestion_errors",
        }
        try:
            found_tables_rows = con.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            ).fetchall()
            found_tables = {row[0] for row in found_tables_rows if not row[0].startswith("sqlite_")}
        except sqlite3.DatabaseError as exc:
            checks.append(
                VerificationCheckResult(
                    check_id="DB-02",
                    category="Database",
                    description="Relational 3NF schema tables presence (10 tables)",
                    expected=f"All 10 tables: {', '.join(sorted(expected_tables))}",
                    observed=f"Corrupt database file ({exc})",
                    passed=False,
                )
            )
            return checks, False

        missing_tables = expected_tables - found_tables
        tables_ok = len(missing_tables) == 0

        checks.append(
            VerificationCheckResult(
                check_id="DB-02",
                category="Database",
                description="Relational 3NF schema tables presence (10 tables)",
                expected=f"All 10 tables: {', '.join(sorted(expected_tables))}",
                observed=(
                    f"All 10 tables confirmed present ({len(found_tables)} total tables)"
                    if tables_ok
                    else f"Missing tables: {', '.join(sorted(missing_tables))}"
                ),
                passed=tables_ok,
            )
        )

        # 3. Minimum Candidates Record Count (>= 10)
        if "candidates" in found_tables:
            cand_count = con.execute("SELECT count(*) FROM candidates").fetchone()[0]
            cand_ok = cand_count >= 10
            observed_cand = f"{cand_count} candidates present"
        else:
            cand_ok = False
            observed_cand = "Table 'candidates' missing"

        checks.append(
            VerificationCheckResult(
                check_id="DB-03",
                category="Data Volume",
                description="Minimum candidate profiles count",
                expected=">= 10 candidate records",
                observed=observed_cand,
                passed=cand_ok,
            )
        )

        # 4. Minimum Job Descriptions Record Count (>= 4)
        if "job_descriptions" in found_tables:
            job_count = con.execute("SELECT count(*) FROM job_descriptions").fetchone()[0]
            job_ok = job_count >= 4
            observed_job = f"{job_count} job vacancies present"
        else:
            job_ok = False
            observed_job = "Table 'job_descriptions' missing"

        checks.append(
            VerificationCheckResult(
                check_id="DB-04",
                category="Data Volume",
                description="Minimum job vacancies count",
                expected=">= 4 job vacancy records",
                observed=observed_job,
                passed=job_ok,
            )
        )

        # 5. Minimum Operational Metrics Record Count (>= 36)
        if "operational_metrics" in found_tables:
            metric_count = con.execute("SELECT count(*) FROM operational_metrics").fetchone()[0]
            metric_ok = metric_count >= 36
            observed_metric = f"{metric_count} operational metric records present"
        else:
            metric_ok = False
            observed_metric = "Table 'operational_metrics' missing"

        checks.append(
            VerificationCheckResult(
                check_id="DB-05",
                category="Data Volume",
                description="Minimum operational hiring metrics count",
                expected=">= 36 operational metric records",
                observed=observed_metric,
                passed=metric_ok,
            )
        )

        # 6. Minimum Matching Evaluations Record Count (>= 40)
        if "matching_evaluations" in found_tables:
            eval_count = con.execute("SELECT count(*) FROM matching_evaluations").fetchone()[0]
            eval_ok = eval_count >= 40
            observed_eval = f"{eval_count} matching evaluations present"
        else:
            eval_ok = False
            observed_eval = "Table 'matching_evaluations' missing"

        checks.append(
            VerificationCheckResult(
                check_id="DB-06",
                category="Data Volume",
                description="Minimum matching evaluations count",
                expected=">= 40 matching evaluation records",
                observed=observed_eval,
                passed=eval_ok,
            )
        )

        # 7. SQLite Foreign Key Integrity Check
        try:
            fk_violations = con.execute("PRAGMA foreign_key_check").fetchall()
            fk_ok = len(fk_violations) == 0
            observed_fk = (
                "0 foreign key violations detected"
                if fk_ok
                else f"{len(fk_violations)} FK violations found: {fk_violations[:3]}"
            )
        except sqlite3.Error as exc:
            fk_ok = False
            observed_fk = f"Error checking foreign keys: {exc}"

        checks.append(
            VerificationCheckResult(
                check_id="DB-07",
                category="Integrity",
                description="Relational foreign key constraints integrity",
                expected="0 foreign key violations",
                observed=observed_fk,
                passed=fk_ok,
            )
        )

        # 8. Evaluation Scores Bounded in [0.0, 100.0]
        if "matching_evaluations" in found_tables:
            score_rows = con.execute(
                """
                SELECT
                    min(overall_score), max(overall_score),
                    min(skills_score), max(skills_score),
                    min(experience_score), max(experience_score),
                    min(semantic_similarity), max(semantic_similarity),
                    count(*)
                FROM matching_evaluations
                """
            ).fetchone()

            if (
                score_rows
                and score_rows[8] > 0
                and all(score_rows[i] is not None for i in range(8))
            ):
                min_ov, max_ov = score_rows[0], score_rows[1]
                min_sk, max_sk = score_rows[2], score_rows[3]
                min_ex, max_ex = score_rows[4], score_rows[5]
                min_sim, max_sim = score_rows[6], score_rows[7]

                scores_ok = (
                    0.0 <= min_ov <= 100.0
                    and 0.0 <= max_ov <= 100.0
                    and 0.0 <= min_sk <= 100.0
                    and 0.0 <= max_sk <= 100.0
                    and 0.0 <= min_ex <= 100.0
                    and 0.0 <= max_ex <= 100.0
                    and 0.0 <= min_sim <= 1.0
                    and 0.0 <= max_sim <= 1.0
                )
                checks.append(
                    VerificationCheckResult(
                        check_id="DB-08",
                        category="Integrity",
                        description="Evaluation score boundaries within [0.0, 100.0]",
                        expected="All scores in [0.0, 100.0], similarity in [0.0, 1.0]",
                        observed=(
                            f"Overall: [{min_ov:.2f}, {max_ov:.2f}], "
                            f"Skills: [{min_sk:.2f}, {max_sk:.2f}], "
                            f"Exp: [{min_ex:.2f}, {max_ex:.2f}], "
                            f"Sim: [{min_sim:.4f}, {max_sim:.4f}]"
                        ),
                        passed=scores_ok,
                    )
                )
            else:
                checks.append(
                    VerificationCheckResult(
                        check_id="DB-08",
                        category="Integrity",
                        description="Evaluation score boundaries within [0.0, 100.0]",
                        expected="Scores in [0.0, 100.0]",
                        observed="No evaluation records found",
                        passed=False,
                    )
                )
        else:
            checks.append(
                VerificationCheckResult(
                    check_id="DB-08",
                    category="Integrity",
                    description="Evaluation score boundaries within [0.0, 100.0]",
                    expected="Scores in [0.0, 100.0]",
                    observed="Table 'matching_evaluations' missing",
                    passed=False,
                )
            )

    finally:
        con.close()

    # 9. Excel Workbook Existence and Readability
    excel_path = target_reports / "talent_bi_report.xlsx"
    if excel_path.exists() and excel_path.stat().st_size > 0:
        try:
            wb = openpyxl.load_workbook(str(excel_path), data_only=True)
            checks.append(
                VerificationCheckResult(
                    check_id="REP-01",
                    category="Deliverables",
                    description="Excel workbook deliverable existence and readability",
                    expected=f"File exists at {excel_path.name} and readable by openpyxl",
                    observed=f"Exists, size: {excel_path.stat().st_size:,} bytes",
                    passed=True,
                )
            )

            # 10. Excel 4 Required Sheets Presence
            expected_sheets = {
                "Executive Dashboard",
                "Candidates Ranking",
                "Gap Analysis",
                "Operational Metrics",
            }
            found_sheets = set(wb.sheetnames)
            missing_sheets = expected_sheets - found_sheets
            sheets_ok = len(missing_sheets) == 0

            checks.append(
                VerificationCheckResult(
                    check_id="REP-02",
                    category="Deliverables",
                    description="Excel workbook structure (all 4 required sheets)",
                    expected="4 sheets: Executive Dashboard, Candidates Ranking, Gap Analysis, Operational Metrics",
                    observed=(
                        f"All 4 sheets verified: {', '.join(wb.sheetnames)}"
                        if sheets_ok
                        else f"Missing sheets: {', '.join(sorted(missing_sheets))}"
                    ),
                    passed=sheets_ok,
                )
            )

            # 11. Excel Sheets Content Non-Empty
            empty_sheets = []
            for sheet_name in expected_sheets:
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    if ws.max_row < 2 or ws.max_column < 2:
                        empty_sheets.append(sheet_name)
            content_ok = len(empty_sheets) == 0

            checks.append(
                VerificationCheckResult(
                    check_id="REP-03",
                    category="Deliverables",
                    description="Excel sheets content non-empty (data rows present)",
                    expected="max_row >= 2 and max_column >= 2 across all 4 sheets",
                    observed=(
                        "All sheets contain structured data rows and columns"
                        if content_ok
                        else f"Empty sheets detected: {', '.join(empty_sheets)}"
                    ),
                    passed=content_ok,
                )
            )
        except Exception as exc:
            checks.append(
                VerificationCheckResult(
                    check_id="REP-01",
                    category="Deliverables",
                    description="Excel workbook deliverable existence and readability",
                    expected="Valid Excel workbook",
                    observed=f"Error reading workbook: {exc}",
                    passed=False,
                )
            )
    else:
        checks.append(
            VerificationCheckResult(
                check_id="REP-01",
                category="Deliverables",
                description="Excel workbook deliverable existence and readability",
                expected=f"File exists at {excel_path.name}",
                observed="File not found or empty",
                passed=False,
            )
        )

    # 12. PDF Executive Dossier Existence and Size (> 5KB)
    pdf_path = target_reports / "executive_evaluation_summary.pdf"
    if pdf_path.exists():
        pdf_size = pdf_path.stat().st_size
        pdf_size_ok = pdf_size > 5120
        checks.append(
            VerificationCheckResult(
                check_id="REP-04",
                category="Deliverables",
                description="PDF executive summary dossier size (> 5 KB)",
                expected="File size > 5,120 bytes (> 5 KB)",
                observed=f"File exists, size: {pdf_size:,} bytes",
                passed=pdf_size_ok,
            )
        )

        # 13. PDF Magic Header (%PDF-)
        try:
            with open(pdf_path, "rb") as f:
                header = f.read(5)
            header_ok = header.startswith(b"%PDF-")
            checks.append(
                VerificationCheckResult(
                    check_id="REP-05",
                    category="Deliverables",
                    description="PDF document magic header signature",
                    expected="Header starts with %PDF-",
                    observed=f"Header: {header.decode('latin1', errors='replace')}",
                    passed=header_ok,
                )
            )
        except Exception as exc:
            checks.append(
                VerificationCheckResult(
                    check_id="REP-05",
                    category="Deliverables",
                    description="PDF document magic header signature",
                    expected="Header starts with %PDF-",
                    observed=f"Error reading header: {exc}",
                    passed=False,
                )
            )

        # 14. PDF Stream Readability and Page Count
        try:
            reader = pypdf.PdfReader(str(pdf_path))
            page_count = len(reader.pages)
            total_chars = sum(len(p.extract_text() or "") for p in reader.pages)
            stream_ok = page_count >= 1 and total_chars > 0
            checks.append(
                VerificationCheckResult(
                    check_id="REP-06",
                    category="Deliverables",
                    description="PDF document structure and valid text stream",
                    expected=">= 1 pages with non-empty extractable text",
                    observed=f"{page_count} pages confirmed, {total_chars:,} text characters extracted",
                    passed=stream_ok,
                )
            )
        except Exception as exc:
            checks.append(
                VerificationCheckResult(
                    check_id="REP-06",
                    category="Deliverables",
                    description="PDF document structure and valid text stream",
                    expected="Valid PDF structure",
                    observed=f"Error extracting text stream: {exc}",
                    passed=False,
                )
            )
    else:
        checks.append(
            VerificationCheckResult(
                check_id="REP-04",
                category="Deliverables",
                description="PDF executive summary dossier size (> 5 KB)",
                expected=f"File exists at {pdf_path.name}",
                observed="File not found",
                passed=False,
            )
        )

    all_passed = all(check.passed for check in checks)
    return checks, all_passed


# --- CLI Commands ---


@app.command("ingest")
def ingest_command(
    data_dir: Path = typer.Option(
        Path("data/raw"),
        "--data-dir",
        "-d",
        help="Path to directory containing raw JSON/CSV data files.",
    ),
    db_path: Path = typer.Option(
        Path("data/output/talent_bi.db"),
        "--db-path",
        "-b",
        help="Path to destination SQLite database file.",
    ),
) -> None:
    """Extract, validate, and load raw talent datasets into relational storage."""
    console.print(
        Panel(
            "[bold cyan]TALENT ANALYTICS PIPELINE - ETL INGESTION[/bold cyan]\n"
            f"Source Directory: {data_dir}\n"
            f"Target Database:  {db_path}",
            box=box.ROUNDED,
            border_style="cyan",
        )
    )

    orchestrator = PipelineOrchestrator()
    result = orchestrator.run_ingestion(data_dir=data_dir, db_path=db_path)

    table = Table(
        title="ETL Ingestion Summary",
        box=box.ROUNDED,
        header_style="bold cyan",
        show_footer=True,
    )
    table.add_column("Dataset / Entity", style="bold", footer="Total")
    table.add_column("Source File")
    table.add_column("Extracted", justify="right", footer=str(result.total_extracted))
    table.add_column("Validated", justify="right", footer=str(result.total_validated))
    table.add_column("Loaded", justify="right", footer=str(result.total_loaded))
    table.add_column("Quarantined", justify="right", footer=str(result.total_rejected))
    table.add_column("Status", justify="center")

    for entity, batch in result.batch_summaries.items():
        status_styled = (
            "[bold green]COMPLETED[/bold green]"
            if batch.status == "completed"
            else "[bold yellow]PARTIAL[/bold yellow]"
            if batch.status == "partial"
            else "[bold red]FAILED[/bold red]"
        )
        table.add_row(
            entity.replace("_", " ").title(),
            batch.source_filename,
            str(batch.records_extracted),
            str(batch.records_validated),
            str(batch.records_loaded),
            str(batch.records_rejected),
            status_styled,
        )

    console.print(table)
    console.print(
        f"Execution duration: [bold]{result.duration_seconds:.2f}s[/bold] | "
        f"Status: {'[bold green]SUCCESS[/bold green]' if result.success else '[bold red]FAILED[/bold red]'}"
    )

    if not result.success:
        raise typer.Exit(code=1)


@app.command("match")
def match_command(
    db_path: Path = typer.Option(
        Path("data/output/talent_bi.db"),
        "--db-path",
        "-b",
        help="Path to SQLite database file.",
    ),
    persist: bool = typer.Option(
        True,
        "--persist/--no-persist",
        help="Whether to persist matching evaluations to database.",
    ),
    top_n: int = typer.Option(
        5,
        "--top-n",
        "-n",
        min=1,
        help="Number of top candidate matches to display.",
    ),
) -> None:
    """Execute semantic candidate matching and multi-factor adequation scoring."""
    console.print(
        Panel(
            "[bold cyan]TALENT ANALYTICS PIPELINE - SEMANTIC MATCHING[/bold cyan]\n"
            f"Database Path: {db_path}\n"
            f"Persistence:   {'Enabled' if persist else 'Disabled'}",
            box=box.ROUNDED,
            border_style="cyan",
        )
    )

    orchestrator = PipelineOrchestrator()
    result = orchestrator.run_matching(db_path=db_path, persist=persist)

    if not result.success:
        console.print("[bold red]Matching engine failed or produced no evaluations.[/bold red]")
        raise typer.Exit(code=1)

    # Display Top N Match Results
    evals_sorted = sorted(result.evaluations, key=lambda e: float(e.overall_score), reverse=True)
    display_evals = evals_sorted[:top_n]

    table = Table(
        title=f"Top {len(display_evals)} Candidate Matching Evaluations",
        box=box.ROUNDED,
        header_style="bold cyan",
    )
    table.add_column("Candidate ID", style="dim")
    table.add_column("Job ID", style="dim")
    table.add_column("Overall Score", justify="right")
    table.add_column("Fit Tier", justify="center")
    table.add_column("Skills Score", justify="right")
    table.add_column("Exp Score", justify="right")
    table.add_column("Semantic Sim", justify="right")

    for e in display_evals:
        score_val = float(e.overall_score)
        score_styled = (
            f"[bold green]{score_val:.1f}%[/bold green]"
            if score_val >= 70.0
            else f"[bold yellow]{score_val:.1f}%[/bold yellow]"
            if score_val >= 50.0
            else f"[bold red]{score_val:.1f}%[/bold red]"
        )
        rec_styled = (
            "[bold green]High Fit[/bold green]"
            if e.recommendation in ("strong_hire", "hire")
            else "[yellow]Moderate Fit[/yellow]"
            if e.recommendation == "consider"
            else "[red]Low Fit[/red]"
        )
        table.add_row(
            e.candidate_id,
            e.job_id,
            score_styled,
            rec_styled,
            f"{float(e.skills_score):.1f}%",
            f"{float(e.experience_score):.1f}%",
            f"{float(e.semantic_similarity):.4f}",
        )

    console.print(table)

    summary_table = Table(box=box.ROUNDED, header_style="bold cyan", show_header=False)
    summary_table.add_column("Metric", style="bold")
    summary_table.add_column("Value")
    summary_table.add_row("Total Pairs Evaluated", str(result.evaluations_count))
    summary_table.add_row("Cohort Mean Score", f"{result.average_score:.2f}%")
    summary_table.add_row("High Fit Count (>= 70%)", str(result.high_fit_count))
    summary_table.add_row("Execution Duration", f"{result.duration_seconds:.2f}s")
    console.print(summary_table)

    console.print("[bold green]Matching evaluation completed successfully.[/bold green]")


@app.command("report")
def report_command(
    db_path: Path = typer.Option(
        Path("data/output/talent_bi.db"),
        "--db-path",
        "-b",
        help="Path to SQLite database file.",
    ),
    output_dir: Path = typer.Option(
        Path("data/output/reports"),
        "--output-dir",
        "-o",
        help="Destination directory for generated reports.",
    ),
) -> None:
    """Compute BI metrics and render multi-tab Excel workbook and executive PDF dossier."""
    console.print(
        Panel(
            "[bold cyan]TALENT ANALYTICS PIPELINE - BI REPORTING & GENERATION[/bold cyan]\n"
            f"Database Source: {db_path}\n"
            f"Output Folder:   {output_dir}",
            box=box.ROUNDED,
            border_style="cyan",
        )
    )

    orchestrator = PipelineOrchestrator()
    result = orchestrator.run_reporting(db_path=db_path, output_dir=output_dir)

    if not result.success:
        console.print("[bold red]Failed to generate reporting deliverables.[/bold red]")
        raise typer.Exit(code=1)

    table = Table(
        title="Generated Reporting Deliverables",
        box=box.ROUNDED,
        header_style="bold cyan",
    )
    table.add_column("Deliverable", style="bold")
    table.add_column("Format")
    table.add_column("File Path")
    table.add_column("Size (Bytes)", justify="right")

    table.add_row(
        "Multi-Tab Operational Workbook",
        "Excel (.xlsx)",
        str(result.excel_path),
        f"{result.excel_path.stat().st_size:,}",
    )
    table.add_row(
        "Executive Evaluation Summary",
        "PDF (.pdf)",
        str(result.pdf_path),
        f"{result.pdf_path.stat().st_size:,}",
    )
    console.print(table)

    if result.kpi_summary:
        kpi = result.kpi_summary
        kpi_table = Table(
            title="Executive Recruitment KPI & ROI Summary",
            box=box.ROUNDED,
            header_style="bold cyan",
        )
        kpi_table.add_column("Metric Indicator", style="bold")
        kpi_table.add_column("Calculated Value", justify="right")

        kpi_table.add_row("Total Candidate Profiles", str(kpi.total_candidates))
        kpi_table.add_row("Open Job Vacancies", str(kpi.total_jobs))
        kpi_table.add_row("Evaluated Pairs", str(kpi.total_evaluations))
        kpi_table.add_row("Cohort Mean Fit Score", f"{kpi.average_matching_score:.1f}%")
        kpi_table.add_row(
            "High Fit Profiles (>= 70%)",
            f"{kpi.high_fit_count} ({kpi.high_fit_rate:.1f}%)",
        )
        kpi_table.add_row(
            "Manual Screening Baseline",
            f"{kpi.roi.total_manual_hours:.1f} hours",
        )
        kpi_table.add_row(
            "Automated Net Time Saved",
            f"[bold green]{kpi.roi.total_hours_saved:.1f} hours[/bold green]",
        )
        kpi_table.add_row(
            "Estimated Financial Value Created",
            f"[bold green]${kpi.roi.total_cost_saved:,.2f}[/bold green]",
        )
        console.print(kpi_table)

    console.print(
        f"Execution duration: [bold]{result.duration_seconds:.2f}s[/bold] | "
        "[bold green]Reporting deliverables generated successfully.[/bold green]"
    )


@app.command("run-all")
def run_all_command(
    data_dir: Path = typer.Option(
        Path("data/raw"),
        "--data-dir",
        "-d",
        help="Path to directory containing raw JSON/CSV data files.",
    ),
    db_path: Path = typer.Option(
        Path("data/output/talent_bi.db"),
        "--db-path",
        "-b",
        help="Path to destination SQLite database file.",
    ),
    output_dir: Path = typer.Option(
        Path("data/output/reports"),
        "--output-dir",
        "-o",
        help="Destination directory for generated reports.",
    ),
) -> None:
    """Execute complete end-to-end pipeline: Ingestion -> Matching -> Persistence -> Reporting."""
    console.print(
        Panel(
            "[bold cyan]TALENT ANALYTICS PIPELINE - COMPLETE END-TO-END EXECUTION[/bold cyan]\n"
            f"Raw Data Directory: {data_dir}\n"
            f"SQLite Database:    {db_path}\n"
            f"Report Directory:   {output_dir}",
            box=box.ROUNDED,
            border_style="cyan",
        )
    )

    orchestrator = PipelineOrchestrator()
    summary = orchestrator.run_all(data_dir=data_dir, db_path=db_path, output_dir=output_dir)

    stage_table = Table(
        title="Pipeline Execution Stage Breakdown",
        box=box.ROUNDED,
        header_style="bold cyan",
    )
    stage_table.add_column("Pipeline Stage", style="bold")
    stage_table.add_column("Primary Metric Output")
    stage_table.add_column("Duration", justify="right")
    stage_table.add_column("Status", justify="center")

    ingest_status = (
        "[bold green]SUCCESS[/bold green]"
        if summary.ingestion.success
        else "[bold red]FAILED[/bold red]"
    )
    stage_table.add_row(
        "1. ETL Ingestion",
        f"{summary.ingestion.total_loaded} records loaded ({summary.ingestion.total_rejected} quarantined)",
        f"{summary.ingestion.duration_seconds:.2f}s",
        ingest_status,
    )

    match_status = (
        "[bold green]SUCCESS[/bold green]"
        if summary.matching.success
        else "[bold red]FAILED[/bold red]"
    )
    stage_table.add_row(
        "2. Semantic Matching",
        f"{summary.matching.evaluations_count} evaluations (avg: {summary.matching.average_score:.1f}%)",
        f"{summary.matching.duration_seconds:.2f}s",
        match_status,
    )

    report_status = (
        "[bold green]SUCCESS[/bold green]"
        if summary.reporting.success
        else "[bold red]FAILED[/bold red]"
    )
    stage_table.add_row(
        "3. BI Reporting",
        "Excel workbook & PDF executive dossier rendered",
        f"{summary.reporting.duration_seconds:.2f}s",
        report_status,
    )

    console.print(stage_table)

    final_status = (
        "[bold green]PIPELINE EXECUTION COMPLETED SUCCESSFULLY[/bold green]"
        if summary.success
        else "[bold red]PIPELINE EXECUTION ENCOUNTERED ERRORS[/bold red]"
    )
    console.print(
        Panel(
            f"{final_status}\n"
            f"Total Elapsed Time: [bold]{summary.total_duration_seconds:.2f}s[/bold]",
            box=box.ROUNDED,
            border_style="green" if summary.success else "red",
        )
    )

    if not summary.success:
        raise typer.Exit(code=1)


@app.command("verify")
def verify_command(
    db_path: Path = typer.Option(
        Path("data/output/talent_bi.db"),
        "--db-path",
        "-b",
        help="Path to SQLite database file.",
    ),
    reports_dir: Path = typer.Option(
        Path("data/output/reports"),
        "--reports-dir",
        "-r",
        help="Directory containing generated reports.",
    ),
) -> None:
    """Audit pipeline integrity, referential constraints, and deliverable artifacts."""
    console.print(
        Panel(
            "[bold cyan]TALENT ANALYTICS PIPELINE - VERIFICATION AUDIT[/bold cyan]\n"
            f"Database Path:    {db_path}\n"
            f"Reports Directory:{reports_dir}",
            box=box.ROUNDED,
            border_style="cyan",
        )
    )

    checks, all_passed = run_pipeline_verification(db_path=db_path, reports_dir=reports_dir)

    table = Table(
        title="Verification Audit Results",
        box=box.ROUNDED,
        header_style="bold cyan",
    )
    table.add_column("ID", style="dim", justify="center")
    table.add_column("Category", style="bold")
    table.add_column("Verification Check")
    table.add_column("Observed Measurement")
    table.add_column("Status", justify="center")

    for c in checks:
        status_styled = (
            "[bold green]PASS[/bold green]"
            if c.passed
            else "[bold red]FAIL[/bold red]"
        )
        table.add_row(
            c.check_id,
            c.category,
            c.description,
            c.observed,
            status_styled,
        )

    console.print(table)

    passed_count = sum(1 for c in checks if c.passed)
    total_count = len(checks)
    console.print(
        f"Verification Summary: [bold]{passed_count}/{total_count}[/bold] checks passed | "
        f"Status: {'[bold green]PASS[/bold green]' if all_passed else '[bold red]FAIL[/bold red]'}"
    )

    if not all_passed:
        raise typer.Exit(code=1)


def main() -> None:
    """Primary execution entry point for CLI application."""
    app()


if __name__ == "__main__":
    main()
